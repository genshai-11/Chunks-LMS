import json
import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(ROOT, '@output')

pkg_files = [
    ('G01', 'G01-21Q-Ecommerce-1.json', 'G01_Green_Focus'),
    ('G02', 'G02-21Q-Ecommerce-2.json', 'G02_Green_Focus'),
    ('R01', 'R01-21Q-Ecommerce-56V-1.json', 'R01_Red_Awareness'),
    ('R02', 'R02-21Q-Ecommerce-56V-2.json', 'R02_Red_Awareness'),
]

packages = []
for code, fname, sheet_title in pkg_files:
    fpath = os.path.join(OUTPUT_DIR, fname)
    with open(fpath, 'r', encoding='utf-8') as fp:
        packages.append((code, json.load(fp), sheet_title))

wb = openpyxl.Workbook()
# Remove default sheet
wb.remove(wb.active)

# Styles
font_title = Font(name='Calibri', size=14, bold=True, color='1E293B')
font_subtitle = Font(name='Calibri', size=10, italic=True, color='64748B')
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_bold = Font(name='Calibri', size=10, bold=True, color='0F172A')
font_regular = Font(name='Calibri', size=10, color='1E293B')
font_mono = Font(name='Consolas', size=9, color='0F172A')

fill_header_green = PatternFill(start_color='059669', end_color='059669', fill_type='solid') # Emerald
fill_header_red = PatternFill(start_color='DC2626', end_color='DC2626', fill_type='solid') # Rose
fill_header_blue = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid') # Blue
fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0'),
)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center')

# -------------------------------------------------------------------------
# Sheet 1: Overview
# -------------------------------------------------------------------------
ws_over = wb.create_sheet(title='Overview')
ws_over.views.sheetView[0].showGridLines = True

ws_over.cell(row=1, column=1, value='CHUNKS LMS • E-COMMERCE MINI TEST SUITE (7x3 Architecture)').font = font_title
ws_over.cell(row=2, column=1, value='Standardized Focus (Green 12V) and Awareness (Red 56V) mini test packages generated from Improv Set 05.').font = font_subtitle

overview_headers = [
    'Package Code', 'Test Type', 'Target CPD', 'Layout', 'Questions', 'Sessions',
    'CVR Curve (Ω)', 'CCI Curve (A)', 'Audio Languages Preset', 'Title'
]
for col_idx, h in enumerate(overview_headers, start=1):
    c = ws_over.cell(row=4, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_header_blue
    c.alignment = align_center

for row_idx, (code, pkg, _) in enumerate(packages, start=5):
    cvr_curve = [s['targetCvrOhm'] for s in pkg['sections']]
    cci_curve = [s['cciAmpe'] for s in pkg['sections']]
    langs = ' · '.join([s.get('sessionLanguage', 'en').upper() for s in pkg['sections']])
    
    values = [
        pkg['packageCode'],
        pkg['testType'],
        f"{pkg['targetCpd']}V",
        pkg.get('sessionLayout', '7x3'),
        pkg['targetQuestions'],
        len(pkg['sections']),
        str(cvr_curve),
        str(cci_curve),
        langs,
        pkg['title'],
    ]
    for col_idx, val in enumerate(values, start=1):
        c = ws_over.cell(row=row_idx, column=col_idx, value=val)
        c.font = font_bold if col_idx in [1, 2, 3] else font_regular
        c.border = thin_border
        c.alignment = align_center if col_idx in [2, 3, 4, 5, 6] else align_left
        if row_idx % 2 == 0:
            c.fill = fill_zebra

# -------------------------------------------------------------------------
# Sheets 2 & 3: Green Tests (G01 & G02)
# -------------------------------------------------------------------------
for code, pkg, sheet_title in packages:
    if pkg['testType'] != 'GREEN':
        continue
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    ws.cell(row=1, column=1, value=f"{pkg['packageCode']} — {pkg['title']}").font = font_title
    ws.cell(row=2, column=1, value='Fluent Complete Sentences (TL = 1.0) • Target CPD 12V • Per-Session TTS Audio').font = font_subtitle

    headers = [
        'Session', 'Item', 'Global #', 'Part', 'Level', 'CVR (Ω)', 'CCI (A)', 'CPD (V)',
        'Term (EN)', 'Term (VI)', 'Complete Sentence (EN)', 'Complete Sentence (VI)',
        'Audio Lang', 'Voice Model'
    ]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_green
        c.alignment = align_center

    cur_row = 5
    for sec in pkg['sections']:
        s_num = sec['sectionOrder']
        part = sec.get('part', 1)
        lang = sec.get('sessionLanguage', 'en')
        voice = sec.get('voiceModel', 'google/en-US-Neural2-F')

        for it in sec['items']:
            vals = [
                s_num,
                it['itemOrder'],
                it['globalItemNumber'],
                f"Part {part}",
                it.get('sentenceLevel', 'A2'),
                it.get('measuredCvr', sec['targetCvrOhm']),
                sec['cciAmpe'],
                it.get('cvrBreakdown', {}).get('cpd', sec.get('cpd', 12)),
                it.get('termEn', ''),
                it.get('termVi', ''),
                it.get('promptEn', ''),
                it.get('promptVi', ''),
                lang.upper(),
                voice,
            ]
            for col_idx, val in enumerate(vals, start=1):
                c = ws.cell(row=cur_row, column=col_idx, value=val)
                c.font = font_bold if col_idx in [1, 2, 3, 6, 7, 8, 13] else font_regular
                c.border = thin_border
                c.alignment = align_center if col_idx in [1, 2, 3, 4, 5, 6, 7, 8, 13] else align_left
                if cur_row % 2 == 0:
                    c.fill = fill_zebra
            cur_row += 1

# -------------------------------------------------------------------------
# Sheets 4 & 5: Red Tests (R01 & R02)
# -------------------------------------------------------------------------
for code, pkg, sheet_title in packages:
    if pkg['testType'] != 'RED':
        continue
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True

    ws.cell(row=1, column=1, value=f"{pkg['packageCode']} — {pkg['title']}").font = font_title
    ws.cell(row=2, column=1, value='Cognitive Traps & 650ms SSML Breaks • Hint Progression [2,3,4,2,3,4,4] • Target CPD 56V').font = font_subtitle

    headers = [
        'Session', 'Item', 'Global #', 'Part', 'Hint Count', 'CVR (Ω)', 'CCI (A)', 'CPD (V)',
        'Hints Sequence (EN)', 'Hints Sequence (VI)', 'Spoken SSML Script (650ms Break)',
        'Audio Lang', 'Voice Model'
    ]
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header_red
        c.alignment = align_center

    cur_row = 5
    for sec in pkg['sections']:
        s_num = sec['sectionOrder']
        part = sec.get('part', 1)
        lang = sec.get('sessionLanguage', 'vi')
        voice = sec.get('voiceModel', 'google/vi-VN-Neural2-A')
        hc = sec.get('hintCount', 2)

        for it in sec['items']:
            spoken = it.get('spokenScriptVi') if lang == 'vi' else it.get('spokenScriptEn')
            if not spoken:
                spoken = it.get('spokenScriptEn') or it.get('spokenScriptVi') or ''

            vals = [
                s_num,
                it['itemOrder'],
                it['globalItemNumber'],
                f"Part {part}",
                f"{hc} hints",
                it.get('measuredCvr', sec['targetCvrOhm']),
                sec['cciAmpe'],
                it.get('cvrBreakdown', {}).get('cpd', sec.get('cpd', 56)),
                it.get('promptEn', ''),
                it.get('promptVi', ''),
                spoken,
                lang.upper(),
                voice,
            ]
            for col_idx, val in enumerate(vals, start=1):
                c = ws.cell(row=cur_row, column=col_idx, value=val)
                c.font = font_bold if col_idx in [1, 2, 3, 5, 6, 7, 8, 12] else (font_mono if col_idx == 11 else font_regular)
                c.border = thin_border
                c.alignment = align_center if col_idx in [1, 2, 3, 4, 5, 6, 7, 8, 12] else align_left
                if cur_row % 2 == 0:
                    c.fill = fill_zebra
            cur_row += 1

# -------------------------------------------------------------------------
# Sheet 6: Master_Catalog (All 84 items)
# -------------------------------------------------------------------------
ws_master = wb.create_sheet(title='Master_Catalog')
ws_master.views.sheetView[0].showGridLines = True
ws_master.cell(row=1, column=1, value='CHUNKS LMS • MASTER CATALOG (All 84 Questions)').font = font_title

master_headers = [
    'Package Code', 'Type', 'Session', 'Item', 'Global #', 'CVR (Ω)', 'CCI (A)', 'CPD (V)',
    'Prompt / Sentence (EN)', 'Prompt / Sentence (VI)', 'Audio Lang', 'Voice Model'
]
for col_idx, h in enumerate(master_headers, start=1):
    c = ws_master.cell(row=3, column=col_idx, value=h)
    c.font = font_header
    c.fill = fill_header_blue
    c.alignment = align_center

m_row = 4
for code, pkg, _ in packages:
    for sec in pkg['sections']:
        s_num = sec['sectionOrder']
        lang = sec.get('sessionLanguage', 'en')
        voice = sec.get('voiceModel', '')
        for it in sec['items']:
            vals = [
                pkg['packageCode'],
                pkg['testType'],
                s_num,
                it['itemOrder'],
                it['globalItemNumber'],
                it.get('measuredCvr', sec['targetCvrOhm']),
                sec['cciAmpe'],
                it.get('cvrBreakdown', {}).get('cpd', sec.get('cpd', 0)),
                it.get('promptEn', ''),
                it.get('promptVi', ''),
                lang.upper(),
                voice,
            ]
            for col_idx, val in enumerate(vals, start=1):
                c = ws_master.cell(row=m_row, column=col_idx, value=val)
                c.font = font_bold if col_idx in [1, 2, 3, 4, 5, 11] else font_regular
                c.border = thin_border
                c.alignment = align_center if col_idx in [2, 3, 4, 5, 6, 7, 8, 11] else align_left
                if m_row % 2 == 0:
                    c.fill = fill_zebra
            m_row += 1

# -------------------------------------------------------------------------
# Auto-adjust column widths for all sheets
# -------------------------------------------------------------------------
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 65)

# Save files
excel_main = os.path.join(OUTPUT_DIR, 'Chunks_LMS_Ecommerce_Package_Tests.xlsx')
excel_alias = os.path.join(OUTPUT_DIR, 'Chunks_LMS_Green_Red_Reused_Improv.xlsx')

wb.save(excel_main)
wb.save(excel_alias)
print(f"✓ Saved {excel_main}")
print(f"✓ Saved {excel_alias}")
